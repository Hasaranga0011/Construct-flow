"""
ConstructFlow Reports Export Generator
Generates PDF, Excel (xlsx), CSV, and multi-project ZIP files.
"""

import io
import csv
import zipfile
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from typing import List, Dict, Any, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

COLOMBO_TZ = ZoneInfo("Asia/Colombo")

def format_lkr(amount: Optional[float]) -> str:
    if amount is None:
        return "N/A"
    return f"Rs. {amount:,.2f}"

def format_date_str(iso_or_date: Optional[str]) -> str:
    if not iso_or_date:
        return "—"
    try:
        clean = iso_or_date.replace("Z", "+00:00")
        if "T" in clean:
            dt = datetime.fromisoformat(clean).astimezone(COLOMBO_TZ)
            return dt.strftime("%d/%m/%Y")
        parts = clean.split("-")
        if len(parts) == 3:
            return f"{parts[2]}/{parts[1]}/{parts[0]}"
        return iso_or_date
    except Exception:
        return str(iso_or_date)

# ---------------------------------------------------------------------------
# Numbered Canvas for PDF Page Numbers
# ---------------------------------------------------------------------------

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#64748B"))
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(self._pagesize[0] - 0.5 * inch, 0.4 * inch, page_text)
        self.drawString(0.5 * inch, 0.4 * inch, "ConstructFlow — Confidential Management Report")
        self.restoreState()


# ---------------------------------------------------------------------------
# PDF Generation
# ---------------------------------------------------------------------------

def generate_pdf_report(
    title: str,
    filters_desc: str,
    project_sections: List[Dict[str, Any]],
    summary_data: Optional[Dict[str, Any]] = None,
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.6 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'RepTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0F172A"),
    )
    subtitle_style = ParagraphStyle(
        'RepSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
    )
    section_head_style = ParagraphStyle(
        'RepSec',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#C2410C"),
    )
    cell_style = ParagraphStyle(
        'RepCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )
    head_cell_style = ParagraphStyle(
        'RepHeadCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
    )

    elements = []

    # Header block
    elements.append(Paragraph("ConstructFlow Management System", subtitle_style))
    elements.append(Spacer(1, 2))
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 4))
    now_str = datetime.now(COLOMBO_TZ).strftime("%d/%m/%Y %I:%M %p")
    header_meta = f"Generated: {now_str} (Asia/Colombo) | Filters: {filters_desc}"
    elements.append(Paragraph(header_meta, subtitle_style))
    elements.append(Spacer(1, 14))

    # Executive Summary table if present
    if summary_data and summary_data.get("metrics"):
        elements.append(Paragraph("Executive Overview", section_head_style))
        elements.append(Spacer(1, 6))
        sum_rows = [[Paragraph(k, cell_style), Paragraph(str(v), cell_style)] for k, v in summary_data["metrics"].items()]
        sum_table = Table([[Paragraph("Indicator", head_cell_style), Paragraph("Value", head_cell_style)]] + sum_rows, colWidths=[200, 200])
        sum_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        elements.append(sum_table)
        elements.append(Spacer(1, 14))

    # Project Sections
    for idx, sec in enumerate(project_sections):
        sec_title = sec.get("project_name", f"Section {idx+1}")
        headers = sec.get("headers", [])
        rows = sec.get("rows", [])
        subtotals = sec.get("subtotals")

        elements.append(Paragraph(sec_title, section_head_style))
        elements.append(Spacer(1, 6))

        if not rows:
            elements.append(Paragraph("No records found matching filters for this section.", subtitle_style))
            elements.append(Spacer(1, 12))
            continue

        table_data = []
        table_data.append([Paragraph(str(h), head_cell_style) for h in headers])

        for r in rows:
            table_data.append([Paragraph(str(c if c is not None else "—"), cell_style) for c in r])

        if subtotals:
            subtotal_cells = []
            for c in subtotals:
                subtotal_cells.append(Paragraph(f"<b>{c}</b>", cell_style))
            table_data.append(subtotal_cells)

        # Dynamic column widths
        col_count = len(headers)
        avail_width = 10.0 * 72.0 # 720pt in landscape
        col_w = max(50.0, avail_width / col_count)

        sec_table = Table(table_data, colWidths=[col_w] * col_count, repeatRows=1)
        t_style = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#C2410C")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ]
        if subtotals:
            t_style.append(('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#FFF7ED")))
            t_style.append(('LINEABOVE', (0, -1), (-1, -1), 1.0, colors.HexColor("#EA580C")))

        sec_table.setStyle(TableStyle(t_style))
        elements.append(sec_table)
        elements.append(Spacer(1, 14))

    doc.build(elements, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Excel (.xlsx) Generation
# ---------------------------------------------------------------------------

def generate_excel_report(
    title: str,
    filters_desc: str,
    project_sections: List[Dict[str, Any]],
    summary_data: Optional[Dict[str, Any]] = None,
) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="C2410C", end_color="C2410C", fill_type="solid")
    subtotal_font = Font(name="Calibri", size=11, bold=True, color="9A3412")
    subtotal_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    # 1. Summary Sheet
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.freeze_panes = "A4"

    ws_sum["A1"] = title
    ws_sum["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    ws_sum["A2"] = f"Filters: {filters_desc} | Generated: {datetime.now(COLOMBO_TZ).strftime('%d/%m/%Y %I:%M %p')}"
    ws_sum["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")

    if summary_data and summary_data.get("metrics"):
        ws_sum["A4"] = "Indicator"
        ws_sum["B4"] = "Value"
        ws_sum["A4"].font = header_font
        ws_sum["A4"].fill = header_fill
        ws_sum["B4"].font = header_font
        ws_sum["B4"].fill = header_fill

        curr_row = 5
        for k, v in summary_data["metrics"].items():
            ws_sum[f"A{curr_row}"] = k
            ws_sum[f"B{curr_row}"] = v
            if isinstance(v, (int, float)):
                ws_sum[f"B{curr_row}"].number_format = "#,##0.00"
            curr_row += 1

    # 2. Project Sheets
    for sec in project_sections:
        raw_name = sec.get("project_name", "Report")
        # Excel sheet name limit 31 chars, forbidden characters : \ / ? * [ ]
        safe_sheet_name = "".join(c for c in raw_name if c not in r":\/?*[]")[:30] or "Project"
        
        # Ensure sheet title uniqueness
        count = 1
        base_name = safe_sheet_name
        while safe_sheet_name in wb.sheetnames:
            safe_sheet_name = f"{base_name[:26]}_{count}"
            count += 1

        ws = wb.create_sheet(title=safe_sheet_name)
        headers = sec.get("headers", [])
        rows = sec.get("rows", [])
        currency_cols = set(sec.get("currency_cols", []))
        numeric_cols = set(sec.get("numeric_cols", []))

        # Title rows
        ws["A1"] = f"{raw_name} — {title}"
        ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="C2410C")
        ws["A2"] = f"Filters: {filters_desc}"
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")

        # Headers on row 4
        ws.freeze_panes = "A5"
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")

        # Data rows
        for row_idx, r in enumerate(rows, 5):
            for col_idx, val in enumerate(r, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_thin
                
                # Numeric or Currency conversion
                if col_idx - 1 in currency_cols or col_idx - 1 in numeric_cols:
                    if isinstance(val, (int, float)):
                        cell.value = float(val)
                    else:
                        try:
                            # Strip Rs. or commas if formatted as string
                            clean_v = str(val).replace("Rs.", "").replace("LKR", "").replace(",", "").strip()
                            cell.value = float(clean_v)
                        except (ValueError, TypeError):
                            cell.value = val

                    if col_idx - 1 in currency_cols:
                        cell.number_format = '"Rs. "#,##0.00'
                    else:
                        cell.number_format = '#,##0.00'
                else:
                    cell.value = val

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# CSV Generation
# ---------------------------------------------------------------------------

def generate_csv_report(
    title: str,
    filters_desc: str,
    project_sections: List[Dict[str, Any]],
) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([f"# ConstructFlow Report: {title}"])
    writer.writerow([f"# Filters: {filters_desc}"])
    writer.writerow([f"# Generated: {datetime.now(COLOMBO_TZ).strftime('%d/%m/%Y %I:%M %p')}"])
    writer.writerow([])

    for sec in project_sections:
        writer.writerow([f"=== {sec.get('project_name', 'Section')} ==="])
        writer.writerow(sec.get("headers", []))
        for r in sec.get("rows", []):
            writer.writerow(r)
        if sec.get("subtotals"):
            writer.writerow(sec.get("subtotals"))
        writer.writerow([])

    return output.getvalue().encode("utf-8")


# ---------------------------------------------------------------------------
# Multi-Project ZIP Generation
# ---------------------------------------------------------------------------

def generate_project_zip(
    title: str,
    filters_desc: str,
    project_sections: List[Dict[str, Any]],
    format_type: str = "pdf",
) -> bytes:
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for idx, sec in enumerate(project_sections):
            p_name = sec.get("project_name", f"project_{idx+1}")
            safe_filename = "".join(c for c in p_name if c.isalnum() or c in (' ', '-', '_')).strip().replace(' ', '_')
            
            if format_type == "xlsx":
                file_bytes = generate_excel_report(title, filters_desc, [sec])
                ext = "xlsx"
            else:
                file_bytes = generate_pdf_report(title, filters_desc, [sec])
                ext = "pdf"

            zf.writestr(f"{safe_filename}_{ext}.{ext}", file_bytes)

    return zip_buffer.getvalue()
